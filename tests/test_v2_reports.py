from io import BytesIO
from copy import deepcopy
import pytest
import openpyxl
from core.models import TerminalState
from core.state import demo_book
from reports.desk_report import generate_desk_report,report_tables
from services.book_risk import book_risk

@pytest.mark.parametrize('section',['all','rates','risk','structured','financing'])
def test_report_sections_use_shared_book_and_have_sources(section):
    state=TerminalState(demo_book('structured'))
    data=generate_desk_report(state,section)
    wb=openpyxl.load_workbook(BytesIO(data),data_only=False)
    assert {'Executive_Summary','Positions','Methodology','Sources'}<=set(wb.sheetnames)
    rows=dict(wb['Executive_Summary'].values)
    assert rows['nav']==pytest.approx(book_risk(state)['nav'])
    assert rows['risk_history_source']=='SYNTHETIC'
    if section in ['all','structured']:assert wb['Structured_Risk'].max_row>1


def test_export_protects_formula_cells_and_does_not_mutate_book():
    state=TerminalState(demo_book());state.book.positions.loc[0,'ticker']='=1+2';before=deepcopy(state.book)
    wb=openpyxl.load_workbook(BytesIO(generate_desk_report(state)))
    assert any(c.value=='=1+2' and c.data_type=='s' for row in wb['Positions'] for c in row)
    assert state.book.positions.equals(before.positions)
    assert report_tables(state)['Risk_Metrics'].set_index('metric').loc['horizon_observations','value']==state.risk.horizon


def test_report_invalid_nav_is_recoverable_in_the_ui():
    from streamlit.testing.v1 import AppTest
    from core.i18n import t

    app = AppTest.from_string('''
import streamlit as st
from core.models import TerminalState
from core.state import demo_book
from components.report_download import render_reports
if 'book_state' not in st.session_state:
    state = TerminalState(demo_book())
    state.book.repo_cash = 1e12
    st.session_state.book_state = state
render_reports(st.session_state.book_state)
''').run()
    app.button[0].click().run()
    assert not app.exception
    assert app.error[0].value == t('book.nav', 'en')
    assert not app.get('download_button')
    app.session_state['book_state'].book.repo_cash = 0.
    app.button[0].click().run()
    assert not app.exception and not app.error
    assert len(app.get('download_button')) == 1
