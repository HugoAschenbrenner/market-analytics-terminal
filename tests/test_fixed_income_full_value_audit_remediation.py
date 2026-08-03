from datetime import date
from io import BytesIO
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from engines.fixed_income_engine import (
    apply_fx_conversion,
    build_currency_exposure_table,
    calculate_bond_risk_metrics,
    calculate_dv01_by_bucket,
    calculate_scenario_pnl,
    estimate_pnl_with_duration_convexity,
    generate_fixed_income_commentary,
    load_bond_data,
    portfolio_summary_to_dict,
    summarize_portfolio,
)
from reports.excel_exporter import generate_fixed_income_risk_report


VALUATION_DATE = date(2025, 4, 1)


def _single_bond() -> pd.DataFrame:
    return pd.DataFrame(
        {
            "bond_id": ["TEST-1"],
            "issuer": ["Test Issuer"],
            "currency": ["EUR"],
            "coupon_rate": [0.06],
            "maturity_date": ["2027-01-01"],
            "issue_date": ["2025-01-01"],
            "frequency": [2],
            "clean_price": [100.0],
            "yield_to_maturity": [0.05],
            "notional": [1_000_000.0],
            "rating": ["A"],
            "sector": ["Corporate"],
            "spread_bps": [100.0],
            "curve_bucket": ["0-2Y"],
        }
    )


def _converted_single_bond() -> pd.DataFrame:
    local = calculate_bond_risk_metrics(
        _single_bond(),
        valuation_date=VALUATION_DATE,
    )
    return apply_fx_conversion(
        local,
        base_currency="EUR",
        fx_rates={"EUR": 1.0},
    )


def test_clean_and_full_values_reconcile_exactly():
    row = _converted_single_bond().iloc[0]

    assert row["clean_market_value"] == pytest.approx(1_000_000.0)
    assert row["full_market_value"] > row["clean_market_value"]
    assert row["full_market_value"] == pytest.approx(
        row["clean_market_value"] + row["accrued_interest_amount"]
    )
    assert row["market_value"] == pytest.approx(row["full_market_value"])


def test_dv01_uses_full_value_not_clean_value():
    row = _converted_single_bond().iloc[0]
    expected_full_dv01 = (
        row["modified_duration"]
        * row["full_market_value"]
        * 0.0001
    )
    clean_value_dv01 = (
        row["modified_duration"]
        * row["clean_market_value"]
        * 0.0001
    )

    assert row["dv01"] == pytest.approx(expected_full_dv01)
    assert row["dv01"] > clean_value_dv01


def test_fx_translation_preserves_clean_full_reconciliation():
    local = calculate_bond_risk_metrics(
        _single_bond(),
        valuation_date=VALUATION_DATE,
    )
    local["currency"] = "USD"

    converted = apply_fx_conversion(
        local,
        base_currency="EUR",
        fx_rates={"USD": 0.92, "EUR": 1.0},
    )
    row = converted.iloc[0]

    assert row["clean_market_value_base"] == pytest.approx(
        row["clean_market_value"] * 0.92
    )
    assert row["full_market_value_base"] == pytest.approx(
        row["full_market_value"] * 0.92
    )
    assert row["full_market_value_base"] == pytest.approx(
        row["clean_market_value_base"]
        + row["accrued_interest_amount_base"]
    )
    assert row["market_value_base"] == pytest.approx(
        row["full_market_value_base"]
    )


def test_summary_uses_full_value_and_reports_clean_value_separately():
    risk_df = _converted_single_bond()
    summary = summarize_portfolio(risk_df, base_currency="EUR")

    assert summary.total_clean_market_value == pytest.approx(
        risk_df["clean_market_value_base"].sum()
    )
    assert summary.total_full_market_value == pytest.approx(
        risk_df["full_market_value_base"].sum()
    )
    assert summary.total_accrued_interest_amount == pytest.approx(
        risk_df["accrued_interest_amount_base"].sum()
    )
    assert summary.total_market_value == pytest.approx(
        summary.total_full_market_value
    )
    assert summary.total_full_market_value > summary.total_clean_market_value


def test_currency_table_and_bucket_table_use_full_value_aliases():
    risk_df = _converted_single_bond()
    currency_df = build_currency_exposure_table(risk_df)
    bucket_df = calculate_dv01_by_bucket(risk_df)

    assert currency_df["market_value_base"].iloc[0] == pytest.approx(
        currency_df["full_market_value_base"].iloc[0]
    )
    assert bucket_df["market_value_base"].iloc[0] == pytest.approx(
        bucket_df["full_market_value_base"].iloc[0]
    )
    assert currency_df["full_market_value_base"].iloc[0] > (
        currency_df["clean_market_value_base"].iloc[0]
    )


def test_parallel_scenario_uses_full_market_value():
    risk_df = _converted_single_bond()
    row = risk_df.iloc[0]
    scenario_df = calculate_scenario_pnl(risk_df)

    actual = scenario_df.loc[
        scenario_df["scenario_name"] == "+25 bps parallel",
        "estimated_pnl",
    ].iloc[0]
    expected = estimate_pnl_with_duration_convexity(
        modified_duration=float(row["modified_duration"]),
        convexity=float(row["convexity"]),
        market_value=float(row["full_market_value_base"]),
        yield_move_bps=25.0,
    )["estimated_pnl"]

    assert actual == pytest.approx(expected)


def test_backward_compatible_synthetic_frame_has_zero_accrual():
    synthetic = pd.DataFrame(
        {
            "bond_id": ["SYN"],
            "currency": ["EUR"],
            "market_value": [1_000_000.0],
            "dv01": [500.0],
        }
    )
    converted = apply_fx_conversion(
        synthetic,
        base_currency="EUR",
        fx_rates={"EUR": 1.0},
    )

    row = converted.iloc[0]
    assert row["clean_market_value"] == pytest.approx(1_000_000.0)
    assert row["full_market_value"] == pytest.approx(1_000_000.0)
    assert row["accrued_interest_amount"] == pytest.approx(0.0)


def test_excel_exposes_clean_full_and_accrued_values():
    bonds = load_bond_data("data/sample_bonds.csv")
    local = calculate_bond_risk_metrics(
        bonds,
        valuation_date=date(2026, 5, 6),
    )
    risk_df = apply_fx_conversion(
        local,
        base_currency="EUR",
        fx_rates={"EUR": 1.0, "USD": 0.92},
    )
    summary = summarize_portfolio(risk_df, base_currency="EUR")
    bucket_df = calculate_dv01_by_bucket(risk_df)
    scenario_df = calculate_scenario_pnl(risk_df)
    currency_df = build_currency_exposure_table(risk_df)
    commentary = generate_fixed_income_commentary(
        risk_df,
        bucket_df,
        scenario_df,
    )

    report_bytes = generate_fixed_income_risk_report(
        summary=portfolio_summary_to_dict(summary),
        risk_df=risk_df,
        bucket_df=bucket_df,
        scenario_df=scenario_df,
        commentary=commentary,
        currency_df=currency_df,
    )
    workbook = load_workbook(BytesIO(report_bytes), read_only=True)

    summary_values = [
        cell.value
        for row in workbook["Summary"].iter_rows()
        for cell in row
        if cell.value is not None
    ]
    bond_headers = [
        cell.value
        for cell in next(workbook["Bond_Level_Risk"].iter_rows())
    ]

    assert "Total Clean Market Value in Base Currency" in summary_values
    assert "Total Full Market Value in Base Currency" in summary_values
    assert "Accrued Interest Amount in Base Currency" in summary_values
    assert "clean_market_value" in bond_headers
    assert "full_market_value" in bond_headers
    assert "accrued_interest_amount" in bond_headers


def test_ui_and_methodology_expose_full_value_contract():
    page = Path("app_pages/fixed_income.py").read_text()
    report = Path("reports/excel_exporter.py").read_text()

    assert "Clean Market Value" in page
    assert "Full Market Value" in page
    assert "Accrued Interest" in page
    assert "DV01 and scenario P&L use full value" in page
    assert "full_market_value_base" in page
    assert "Total Full Market Value in Base Currency" in report
    assert "DV01 and duration/convexity scenario P&L use full market value" in report
