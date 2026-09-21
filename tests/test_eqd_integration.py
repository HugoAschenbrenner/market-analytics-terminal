from dataclasses import replace
from io import BytesIO
from pathlib import Path
import numpy as np
import pandas as pd
import pytest
from streamlit.testing.v1 import AppTest
from services.lab import LabState,option_outputs,curve_outputs
from reports.lab_report import generate_lab_report,lab_report_tables
from engines.scenario_engine import MarketScenario


def app_at(page,version='v2',lang='en'):
    app=AppTest.from_file(str(Path('app.py').resolve()),default_timeout=40)
    app.query_params.update(page=page,version=version,lang=lang,theme='light' if lang=='fr' else 'dark')
    return app


@pytest.mark.parametrize('version',['v1','v2'])
@pytest.mark.parametrize('lang',['en','fr'])
@pytest.mark.parametrize('view',['chain','position','scenario','hedge','greeks'])
def test_eqd_focused_views_render_offline(version,lang,view):
    app=app_at('equity-derivatives',version,lang)
    app.session_state['eqd_view']=view
    app.run()
    assert not app.exception and not app.error
    assert app.get('plotly_chart')
    assert not any('Monte Carlo' in item.value for item in app.title)


@pytest.mark.parametrize('view',['surface','smile','term'])
def test_chain_visualizations_render(view):
    app=app_at('equity-derivatives');app.session_state['eqd_surface_view']=view;app.run()
    assert not app.exception and not app.error
    assert app.get('plotly_chart')


def test_v1_navigation_preserves_lab_inputs_and_dashboard_recomputes():
    app=app_at('equity-derivatives','v1');app.session_state['eqd_view']='scenario';app.run()
    app.number_input(key='eqd_pos_quantity').set_value(-23.).run()
    app.number_input(key='lab_spot_shock').set_value(-8.).run()
    expected=option_outputs(app.session_state.analytics_lab)['scenario']['full']
    app.button(key='v1-nav-cross-asset-dashboard').click().run()
    assert not app.exception and not app.error
    assert app.query_params['page']==['cross-asset-dashboard']
    actual=next(m.value for m in app.metric if m.label=='Option scenario P&L')
    assert actual==f'{expected:,.2f}'
    app.button(key='v1-nav-equity-derivatives').click().run()
    assert app.session_state.analytics_lab.position.quantity==-23
    assert app.session_state.analytics_lab.scenario.equity==-.08


def test_v2_structured_products_has_dedicated_route():
    app=app_at('structured-products').run()
    assert not app.exception and not app.error
    assert app.session_state.terminal.ui.page=='structured-products'
    assert not any(item.key=='eqd_view' for item in app.button_group)
    app.button(key='load_structured_demo').click().run()
    assert not app.exception and not app.error
    assert app.get('plotly_chart')


def test_report_includes_computed_inputs_and_no_formula_injection():
    lab=LabState();lab.chain['underlying']='=2+2';lab.matrix_spots=(-.2,0,.2)
    lab.hedge_scenario=MarketScenario(equity=.1,volatility=.07)
    tables=lab_report_tables(lab);out=option_outputs(lab)
    assert tables['Spot_Vol_Matrix'].spot_shock.unique()==pytest.approx([-.2,0,.2])
    assert tables['Delta_Hedge'].set_index('metric').loc['net_pnl','value']==pytest.approx(out['hedge']['net_pnl'])
    binary=generate_lab_report(lab)
    from openpyxl import load_workbook
    workbook=load_workbook(BytesIO(binary))
    required=['Option_Chain','Cash_Greeks','Scenario_PnL','Delta_Hedge','Zero_Curve','Curve_DV01','Methodology','Sources']
    assert set(required).issubset(workbook.sheetnames)
    assert workbook['Option_Chain']['A2'].value=='=2+2'
    assert workbook['Option_Chain']['A2'].data_type=='s'


def test_curve_builder_and_report_button_render():
    app=AppTest.from_string('from components.curve_builder import render_curve_builder\nrender_curve_builder()',default_timeout=40).run()
    assert not app.exception and not app.error
    app.selectbox(key='lab_curve_scenario').set_value('Bear steepener').run()
    app.button(key='lab_prepare_export').click().run()
    assert not app.exception and not app.error
    assert app.get('download_button')
    assert app.session_state.analytics_lab.scenario.curve_twist==((2.,10.),(10.,40.))


def test_invalid_option_scenario_is_visible_and_recoverable():
    app=app_at('equity-derivatives');app.session_state['eqd_view']='scenario';app.run()
    app.number_input(key='lab_days').set_value(1000.).run()
    assert not app.exception and any('before expiry' in item.value for item in app.warning)
    app.number_input(key='lab_days').set_value(0.).run()
    assert not app.exception and not app.error and app.get('plotly_chart')


def test_zero_risk_attribution_and_export_remain_usable():
    source='''import pandas as pd
from components.risk_attribution import render_attribution
render_attribution(pd.DataFrame({'A':[0.,0.,0.],'B':[0.,0.,0.]}),[.5,.5])'''
    app=AppTest.from_string(source,default_timeout=40).run()
    assert not app.exception and not app.error
    tables=app.session_state.analytics_lab.risk_tables
    assert np.isfinite(tables['Risk_Attribution'].select_dtypes('number')).all().all()
    app.button(key='lab_prepare_export').click().run()
    assert app.get('download_button') and not app.exception
